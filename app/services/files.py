from io import BytesIO
from pathlib import Path

async def extract_text(filename: str, data: bytes) -> str:
    ext=Path(filename).suffix.lower()
    if ext in {'.txt','.md','.py','.json','.csv','.log'}:
        return data.decode('utf-8','replace')[:50000]
    if ext == '.pdf':
        from pypdf import PdfReader
        reader=PdfReader(BytesIO(data))
        return '\n\n'.join((p.extract_text() or '') for p in reader.pages)[:50000]
    if ext == '.docx':
        from docx import Document
        doc=Document(BytesIO(data))
        return '\n'.join(p.text for p in doc.paragraphs)[:50000]
    if ext == '.xlsx':
        import openpyxl
        wb=openpyxl.load_workbook(BytesIO(data),read_only=True,data_only=True)
        chunks=[]
        for ws in wb.worksheets:
            chunks.append(f'SHEET: {ws.title}')
            for row in ws.iter_rows(values_only=True): chunks.append(' | '.join('' if v is None else str(v) for v in row))
        return '\n'.join(chunks)[:50000]
    return ''
